#!/usr/bin/env python3
"""Build Polza_test.xlsx from the CSV and Markdown deliverables (no formulas, text data only).

Sheets: «1-2 База+персонализация», «3 Цепочка», «4 Ваша база», «Ответы (сверх ТЗ)»,
«План запуска (сверх ТЗ)», «Резерв», «README».

Sources (only the prose of the notes and of the README sheet is written here):
  task1_2_enriched.csv   sheet 1-2: the leads; the first six columns are the ones the Polza course asks for
  task1_reserve.csv      sheet «Резерв»: companies of the first version of the base, a department mailbox only
  task1_leads.csv        the leads selected for the base, one per company (the funnel on the README sheet)
  task3_chain.csv        sheet 3; the examples are rendered by tools/export_campaign.py
  task4_final.csv        sheet 4, plus the script verdicts from task4_*_output.csv
  reply_playbook.csv/md  sheet «Ответы»: the reply types and the table of section 4
  launch_plan.md         sheet «План запуска»: the four tables, read by md_tables()

Every number of the README sheet is counted from these files at build time. The same sentences
(number_lines) must stand in README.md, and a few of the numbers in the hand-in notes (HANDIN_NOTES,
a local file that is not part of the repository): after the workbook is saved the script compares them and
returns code 1 if a document is out of date.

Whose mailbox a lead has is counted from the column `тип_адреса`: a personal mailbox, a role mailbox
printed in the person's card (dir@, kd@) and a personal mailbox on a free mail domain are three different
things, and the texts name them separately. A lead whose page is closed in the site's robots.txt must
not be in the base: the build stops on such a row. A lead whose page did not answer on the last
enrichment stays, and both README texts name its site.

No network, nothing is sent.

Usage (from the project root):
  .venv/bin/python tools/build_xlsx.py            build the workbook and check the documents (code 1: out of date)
  .venv/bin/python tools/build_xlsx.py --numbers  print the number sentences only, write nothing
"""

from __future__ import annotations

import csv
import math
import re
import sys
from collections import Counter
from pathlib import Path
from types import SimpleNamespace
from urllib.parse import urlparse

from check_chain import GREETING_COLUMN
from enrich_base import FREE_PERSONAL_BOX, PERSONAL_BOX, SERVICE_BOX
from export_campaign import (
    CAMPAIGN_BY_VERTICAL,
    CAMPAIGN_THEIR,
    LATER_BATCHES,
    NAMED_VARIANT,
    OWN,
    OWN_CAMPAIGNS,
    RESERVE,
    ROLE_BOX_NOTE,
    SUBJECT_A,
    SUBJECT_B,
    TEST_BATCHES,
    THEIR,
    THEIR_BATCH,
    THEIR_NOTE,
    VARIANT_NOTES,
    build_contacts,
    load_chain,
    preview_examples,
    render_chain,
)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "Polza_test.xlsx"
LAUNCH_MD = ROOT / "launch_plan.md"
REPLY_MD = ROOT / "reply_playbook.md"
README_MD, HANDIN_NOTES = "README.md", "SUBMISSION.md"  # the second one is checked only where it exists

SHEET_BASE, SHEET_CHAIN, SHEET_THEIR = "1-2 База+персонализация", "3 Цепочка", "4 Ваша база"
SHEET_REPLIES, SHEET_LAUNCH, SHEET_README = "Ответы (сверх ТЗ)", "План запуска (сверх ТЗ)", "README"
SHEET_RESERVE = "Резерв"
# sheet 3: the column `приветствие_без_имени` of task3_chain.csv
GREETING_HEADER = "Первая строка для запасных вариантов «а» и «б»"

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
SMALL = Font(name=FONT, size=9, italic=True)
TITLE = Font(name=FONT, size=13, bold=True)
SECTION = Font(name=FONT, size=11, bold=True, color="1F3864")
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FILL_OK = PatternFill("solid", fgColor="E2EFDA")
FILL_BAD = PatternFill("solid", fgColor="FCE4D6")
FILL_REVIEW = PatternFill("solid", fgColor="FFF2CC")

# Sheet 1-2: (column of task1_2_enriched.csv, header, width).
BASE_COLUMNS = [
    # the six columns of the course (lesson 2, item 4.4), in its order
    ("Имя", "Имя", 14), ("Фамилия", "Фамилия", 16), ("Должность", "Должность", 28), ("Email", "Email", 28),
    ("Телефон", "Телефон", 20), ("Компания", "Компания", 26),
    # the rest of the name and the company
    ("Отчество", "Отчество", 16), ("компания_в_письме", "{{companyName}} в письме", 20), ("site", "Сайт", 24),
    # the decision maker and how email 1 addresses the company
    ("уровень_контакта", "Уровень контакта", 11), ("вариант_письма_1", "Вариант письма 1", 11),
    ("обращаться_по_имени", "Обращаться по имени", 13), ("должность_в_письме", "{{jobTitle}} в письме", 24),
    ("сегмент_ЛПР", "Сегмент ЛПР", 16), ("источник_имени", "Источник имени", 34),
    ("дата_источника_имени", "Дата источника имени", 16),
    # the address and its validation
    ("email_отдела", "Email отдела (запасной)", 26), ("contact_role", "Контакт в базе задания 1", 28),
    ("email_source", "Где опубликован email", 30), ("тип_адреса", "Тип адреса", 16),
    ("валидация", "Валидация", 40), ("дата_проверки", "Дата проверки", 12),
    # segments and triggers
    ("вертикаль", "Вертикаль", 20), ("segment", "Сегмент", 30),
    ("sales_signal", "Сигнал: есть продажи / ищут клиентов", 46), ("тип_триггера", "Тип триггера", 18),
    ("приоритет", "Приоритет", 10),
    # geography
    ("город", "Город", 16), ("регион_группа", "Регион", 16), ("часовой_пояс", "Часовой пояс", 10),
    # personalisation and the pain hypothesis
    ("Персонализация", "Персонализация", 60), ("Источник", "Источник факта", 34),
    ("Гипотеза_боли", "Гипотеза_боли", 50), ("Проверка_соответствия", "Проверка_соответствия", 18),
]
BASE_TAIL = [("Комментарий", "Комментарий: вычитка и скрипт", 80), ("примечание", "Примечание", 50)]
COURSE_HEADERS = ["Имя", "Фамилия", "Должность", "Email", "Телефон", "Компания"]
RESERVE_REASON = ("причина_резерва", "Почему в резерве", 44)  # the extra first column of task1_reserve.csv

# Sheet 4: (column of task4_final.csv, header, width).
THEIR_COLUMNS = [
    ("company", "company", 22), ("компания_в_письме", "{{companyName}} в письме", 18), ("email", "email", 26),
    ("site", "site", 20), ("проверка", "проверка", 14), ("что_не_так", "что_не_так", 70),
    ("исправленный_сайт", "исправленный_сайт", 26), ("исправленный_email", "исправленный_email", 30),
    ("персонализация", "персонализация", 50), ("источник", "источник", 34), ("сегмент", "сегмент", 26),
    ("язык_письма", "язык_письма", 10), ("триггер", "триггер", 24), ("Гипотеза_боли", "Гипотеза_боли", 50),
]

# Sheet «Ответы»: (column of reply_playbook.csv, header, width).
REPLY_COLUMNS = [
    ("код", "Код", 8), ("тип_по_курсу", "Тип по курсу", 18), ("пример_входящего", "Пример входящего", 40),
    ("срок_ответа", "Срок ответа", 24), ("шаблон_номер", "Шаблон", 14), ("шаблон_текст", "Текст шаблона", 80),
    ("статус_лида", "Статус лида", 24), ("цепочка", "Цепочка", 24), ("следующий_шаг", "Следующий шаг", 50),
]
REPLY_FIELDS_SECTION = "4."  # reply_playbook.md: «4. Что фиксируем»
LAUNCH_TABLES = 4            # launch_plan.md: checklist, calendar, metrics, scaling
LAUNCH_WIDTHS = [30, 52, 44, 32, 30, 34]

# Orders and names for the number sentences.
LEVELS, VARIANTS, PRIORITIES = "АБВ", "вба", "ABC"
TRIGGERS = ["нанимают в продажи", "набирают дилеров и партнёров", "расширение", "выставка", "события нет"]
LPR_ROLES = ["продажи", "первое лицо", "развитие", "маркетинг"]
REGIONS = ["Москва и МО", "Санкт-Петербург", "регионы"]
ZONES = ["МСК", "МСК+1", "МСК+2", "МСК+4", "МСК+6"]
VERTICAL_NAMES = {
    "B2B SaaS": "B2B SaaS", "Интеграторы и автоматизация": "интеграторы и автоматизация", "HR-tech": "HR-tech",
    "Логистика": "складская логистика", "B2B-маркетинг": "B2B-маркетинг", "Упаковка": "упаковка",
    "Промоборудование": "промышленное оборудование", "Промдистрибуция": "промышленная дистрибуция",
    "Юридические услуги": "юридические услуги", "Аудит и консалтинг": "аудит и консалтинг",
    "Подбор руководителей": "подбор руководителей", "Логистика ВЭД": "логистика ВЭД",
    "Инжиниринг и промбезопасность": "инжиниринг и промбезопасность", "Оптовая дистрибуция": "оптовая дистрибуция",
    "Стройматериалы": "стройматериалы", "Коммерческая техника": "коммерческая техника",
    "Медицинские изделия": "медицинские изделия", "Полиграфия и сувениры": "полиграфия и сувениры",
}
# A lead's address by `тип_адреса`: type -> how the texts call it. Every row of the base has one of these.
LEAD_TYPES = {PERSONAL_BOX: "именных адресов", SERVICE_BOX: "ящик должности ЛПР",
              FREE_PERSONAL_BOX: "личный ящик на бесплатном домене"}
# What a lead is; README.md and the README sheet say it in these words.
LEAD_RULE = ("Лид — это ЛПР (первое лицо, коммерческий директор, руководитель продаж, маркетинга или развития), чей "
             "рабочий адрес компания сама напечатала на своём сайте рядом с его именем: в одной карточке, строке или "
             "блоке, не дальше 350 знаков. Ящик должности из карточки человека (dir@, kd@) считается, но помечен; "
             "общие ящики (info@, sales@, zakaz@, office@) лидом не считаются.")
# The research behind the base (working files that are not part of the repository): candidates checked by the
# second pass, rejected by it, removed by the final check (two contact pages closed in the site's robots.txt, a job
# title the site does not confirm, an address whose owner is unclear, a name the company page still prints although
# a trade publication already calls the person a former director), spare contacts of a company that already
# has a lead, leads held back by hand.
RESEARCH = SimpleNamespace(candidates=126, rejected=13, final=5, spare=8, held=2)
FINAL_REASONS = ("у двух страница с контактом закрыта в robots.txt сайта, у одного сайт не подтверждает должность, "
                 "у одного неясно, чей адрес напечатан в карточке, у одного имя на странице устарело: отраслевое "
                 "издание называет этого человека уже бывшим директором")
# The headline numbers of LEADFINDER.md; README.md repeats the sentence, LEADFINDER.md holds the measurements.
LEADFINDER_LINE = ("На размеченных вручную страницах (99 строк, 90 сайтов) leadfinder находит 77 из 88 подтверждённых "
                   "лидов (87,5 %); среди 132 выданных строк неверных нет, три — с оговоркой; на холодном списке лид "
                   "находится у 2–5 компаний из ста.")
PAST_TENSE_MARK = "прошедшее время"  # `примечание` of a row whose fact is an event that is still ahead
MOBILE_CODE_RE = re.compile(r"^\D*[78]?\D*9\d\d\D")  # a number on a mobile code (9xx), whoever it belongs to
RESERVE_GROUPS = (  # start of `причина_резерва` -> how the texts call the group
    ("ЛПР назван на сайте, но прямого адреса", "ЛПР назван, но адреса рядом с именем нет"),
    ("ЛПР на сайте не назван", "ЛПР на сайте не назван"),
)
RESERVE_OTHER = "в карточке ЛПР общий ящик или адрес, который в базу не записывается"
NO_ANSWER_MARK = "не перепроверен"  # `примечание` of a row whose page did not answer on the last enrich_base.py run
README_NO_ANSWER = "не ответила страница с контактом"  # README.md, limitations: the paragraph that names those sites
LANGUAGE_COLUMN = "язык_письма"   # named in the reason why export_campaign.py leaves a non-Russian row out
NO_DATA = "нет данных"            # the personalisation of a row the script could not do
ROBOTS_MARK = "robots.txt"        # `примечание` of a lead whose page is closed to crawlers (tools/enrich_base.py)


# ---------- reading ----------

def read_csv(name: str) -> list[dict[str, str]]:
    with (ROOT / name).open(encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def plain(text: str) -> str:
    """A Markdown table cell as plain text: no escaped pipes, bold marks or backticks."""
    return text.replace("\\|", "|").replace("**", "").replace("`", "").strip()


def md_tables(path: Path) -> list[tuple[str, list[str], list[list[str]]]]:
    """Tables of a Markdown file as (the «## » section heading above the table, header cells, data rows)."""
    tables, heading, block = [], "", []
    for line in [*path.read_text(encoding="utf-8").splitlines(), ""]:
        if line.startswith("|"):
            block.append(line)
            continue
        if block:
            # split on pipes that are not escaped; the second line is the |---| separator
            cells = [[plain(cell) for cell in re.split(r"(?<!\\)\|", row.strip())[1:-1]] for row in block]
            tables.append((heading, cells[0], cells[2:]))
            block = []
        if line.startswith("## "):  # sections only: «### 4. …» of another section must not pass for «## 4. …»
            heading = line.removeprefix("## ").strip()
    return tables


def first_line(path: Path) -> str:
    """The disclaimer both plans start with («Это план. … ничего не отправлялось …»)."""
    return path.read_text(encoding="utf-8").splitlines()[0].strip()


# ---------- numbers ----------

def plural(number: int, one: str, few: str, many: str) -> str:
    """Russian plural form for a count: 1 строка, 3 строки, 5 строк."""
    if number % 10 == 1 and number % 100 != 11:
        return one
    if number % 10 in (2, 3, 4) and number % 100 not in (12, 13, 14):
        return few
    return many


def listing(counts: Counter, order, names: dict[str, str] | None = None, quote: bool = False) -> str:
    """«А — 5, Б — 34, В — 17»: the keys of `order` that occur, then any other key."""
    keys = [key for key in order if counts.get(key)] + [key for key in counts if key not in order]
    label = (lambda key: f"«{key}»") if quote else (lambda key: (names or {}).get(key, key))
    return ", ".join(f"{label(key)} — {counts[key]}" for key in keys)


def review_status(comment: str) -> str:
    """Review outcome from the note at the start of the comment."""
    if comment.startswith("вычитка: заменено"):
        return "заменено другим фактом"
    if comment.startswith("вычитка: факт скрипта"):
        return "факт скрипта, поправлена формулировка"
    if comment.startswith("вычитка: сверено"):
        return "оставлено как у скрипта"
    return ""


def collect_stats(rows: list[dict], their: list[dict], contacts: list[dict], report, chain_rows: list[dict],
                  launch: list, replies: list[dict], reserve: list[dict], leads_before: int,
                  script_rows: list[dict]) -> SimpleNamespace:
    """Everything the README sheet and the documents say in numbers, counted from the files."""
    def count(field: str, source: list[dict] = rows) -> Counter:
        return Counter(row[field].strip() for row in source)

    def brands(source: list[dict]) -> list[str]:
        return [row["компания_в_письме"] for row in source]

    not_leads = [row["Компания"] for row in rows
                 if row["тип_адреса"] not in LEAD_TYPES or row["уровень_контакта"] != "А" or not row["Имя"]
                 or row["обращаться_по_имени"] != "да"]
    if not_leads:
        raise SystemExit("task1_2_enriched.csv: база состоит из лидов (уровень А, имя, адрес из карточки ЛПР), а эти "
                         f"строки — нет: {', '.join(not_leads)}")
    if RESEARCH.candidates - RESEARCH.rejected - RESEARCH.final - RESEARCH.spare - RESEARCH.held != leads_before:
        raise SystemExit(f"воронка не сходится: {RESEARCH} и {leads_before} строк в task1_leads.csv")
    if leads_before != len(rows):
        raise SystemExit(f"в task1_leads.csv {leads_before} строк, в базе {len(rows)}: пересоберите базу "
                         "(tools/build_base_all.py) и персонализацию недостающих строк")
    closed = [row["Компания"] for row in rows if ROBOTS_MARK in row["примечание"]]
    if closed:
        raise SystemExit("task1_2_enriched.csv: у этих строк страница закрыта в robots.txt сайта, в базу такие строки "
                         f"не идут: {', '.join(closed)}")
    region_only = [row for row in rows if row["город"].endswith("обл.")]
    verticals = count("вертикаль")
    own = [c for c in contacts if c["base"] == OWN]
    in_test = [c for c in own if c["subject_variant"] in (SUBJECT_A, SUBJECT_B)]
    review = Counter(review_status(row["Комментарий"]) for row in rows)
    unknown = sorted(set(verticals) - set(CAMPAIGN_BY_VERTICAL))
    if unknown:
        raise SystemExit(f"вертикали {unknown} нет в CAMPAIGN_BY_VERTICAL (tools/export_campaign.py)")
    reasons = Counter()
    for row in reserve:
        label = next((label for start, label in RESERVE_GROUPS if row["причина_резерва"].startswith(start)), "")
        reasons[label or RESERVE_OTHER] += 1
    words = [len(row["Персонализация"].split()) for row in rows]
    script = Counter(row["Проверка_соответствия"].split(":")[0] for row in script_rows)
    script_empty = sum(1 for row in script_rows if row["Персонализация"].strip().lower() in ("", NO_DATA))
    return SimpleNamespace(
        script_ok=script["OK"], script_other=len(script_rows) - script["OK"], script_empty=script_empty,
        script_done=len(script_rows) - script_empty,
        total=len(rows), leads_before=leads_before,
        kinds=count("тип_адреса"), personal=count("тип_адреса")[PERSONAL_BOX],
        role_names=brands([row for row in rows if row["тип_адреса"] == SERVICE_BOX]),
        patronymic=sum(1 for row in rows if row["Отчество"]),
        phones=sum(1 for row in rows if row["Телефон"]),
        mobile_code=brands([row for row in rows if MOBILE_CODE_RE.match(row["Телефон"])]),
        levels=count("уровень_контакта"), variants=count("вариант_письма_1"), priorities=count("приоритет"),
        triggers=count("тип_триггера"), regions=count("регион_группа"), zones=count("часовой_пояс"),
        roles=count("сегмент_ЛПР"), verticals=verticals,
        by_campaign={campaign: sum(n for v, n in verticals.items() if CAMPAIGN_BY_VERTICAL[v] == campaign)
                     for campaign in OWN_CAMPAIGNS},
        cities=len({row["город"] for row in rows if row not in region_only}), region_only=len(region_only),
        first_wave=[row["компания_в_письме"] for row in rows if row["приоритет"] == "A"],
        last_wave=[row["компания_в_письме"] for row in rows if row["приоритет"] == "C"],
        no_answer=[urlparse(row["site"]).netloc.removeprefix("www.") for row in rows
                   if NO_ANSWER_MARK in row["примечание"]],
        future_facts=brands([row for row in rows if PAST_TENSE_MARK in row["примечание"]]),
        checked=sorted({row["дата_проверки"] for row in rows}),
        review_kept=review["оставлено как у скрипта"], review_reworded=review["факт скрипта, поправлена формулировка"],
        review_replaced=review["заменено другим фактом"], words_min=min(words), words_max=max(words),
        reserve=len(reserve), reserve_reasons=reasons,
        their_total=len(their), their_segments=count("сегмент", their), their_languages=count("язык_письма", their),
        their_swapped=sum(1 for row in their if row["проверка"].startswith("ПЕРЕПУТАНО")),
        all_rows=report.own_rows + report.their_rows, import_rows=len(contacts),
        excluded_language=sum(1 for *_, reason in report.excluded if LANGUAGE_COLUMN in reason),
        excluded_other=sum(1 for *_, reason in report.excluded if LANGUAGE_COLUMN not in reason),
        letters=len(contacts) * (1 + len({row["шаг"] for row in chain_rows if row["шаг"].strip() != "1"})),
        test_rows=len(in_test), test=Counter(c["subject_variant"] for c in in_test),
        waiting=sum(1 for c in own if c not in in_test),
        campaigns=Counter(c["campaign"] for c in contacts),
        max_words=max(int(row["слов_тема_плюс_тело_макс"]) for row in chain_rows),
        checklist_steps=len(launch[0][2]), launch_rows=sum(len(table[2]) for table in launch),
        reply_types=len(replies),
    )


def counted(number: int, names: list[str]) -> str:
    """«2 (Интерволга, TEAMLY)»; just «0» when there is nobody to name."""
    return f"{number} ({', '.join(names)})" if names else str(number)


def number_lines(s: SimpleNamespace) -> dict[str, str]:
    """The sentences with numbers. The README sheet shows them, README.md must contain each one verbatim."""
    rows_word = plural(s.import_rows, "строка", "строки", "строк")
    by_campaign = []
    for campaign in OWN_CAMPAIGNS:
        order = [vertical for vertical, name in CAMPAIGN_BY_VERTICAL.items() if name == campaign]
        inside = listing(Counter({v: s.verticals[v] for v in order if s.verticals[v]}), order, VERTICAL_NAMES)
        by_campaign.append(f"{campaign} — {s.by_campaign[campaign]}: {inside}")
    region_only = ""
    if s.region_only:
        companies = plural(s.region_only, "компании", "компаний", "компаний")
        region_only = f" (ещё у {s.region_only} {companies} указана только область)"
    kinds = ", ".join(
        f"{label} — " + (counted(s.kinds[kind], s.role_names) if kind == SERVICE_BOX else str(s.kinds[kind]))
        for kind, label in LEAD_TYPES.items())
    reserve = ", ".join(f"{label} — {number}" for label, number in s.reserve_reasons.most_common())
    return {
        "leads": (f"Лидов в базе — {s.total}, по одному на компанию: {kinds}. Отчество на сайте есть у "
                  f"{s.patronymic}, телефон отдела или общий — у {s.phones}."),
        "roles": f"Роли ЛПР: {listing(s.roles, LPR_ROLES)}.",
        "verticals": f"Вертикалей — {len(s.verticals)}, в трёх группах: " + "; ".join(by_campaign) + ".",
        "geo": (f"Города: {s.cities}{region_only}. {listing(s.regions, REGIONS)}. "
                f"Часовые пояса: {listing(s.zones, ZONES)}."),
        "levels": (f"Уровень контакта: {listing(s.levels, LEVELS)}. "
                   f"Вариант письма 1: {listing(s.variants, VARIANTS, quote=True)}."),
        "triggers": f"Триггеры: {listing(s.triggers, TRIGGERS)}.",
        "priority": f"Приоритет: {listing(s.priorities, PRIORITIES)}.",
        "export": (f"В файле импорта {s.import_rows} {rows_word} из "
                   f"{s.all_rows}, писем в цепочке — {s.letters}. A/B темы письма 1: {s.test_rows} "
                   f"{plural(s.test_rows, 'строка', 'строки', 'строк')}, {listing(s.test, [SUBJECT_A, SUBJECT_B])}. "
                   f"Кампании: {listing(s.campaigns, [*OWN_CAMPAIGNS, CAMPAIGN_THEIR])}."),
        "reserve": f"Резерв — {s.reserve} {plural(s.reserve, 'компания', 'компании', 'компаний')}: {reserve}.",
        "their": (f"Ваша база: {listing(s.their_segments, [])}. "
                  f"Язык письма: {listing(s.their_languages, ['RU', 'EN'])}."),
    }


def funnel_line(s: SimpleNamespace) -> str:
    """How the base was narrowed down; README.md must contain it verbatim."""
    passed = RESEARCH.candidates - RESEARCH.rejected
    return (f"Воронка: {RESEARCH.candidates} кандидатов (ЛПР с адресом на сайте компании) → вторая, независимая "
            f"проверка по живой странице подтвердила {passed} и отклонила {RESEARCH.rejected} → финальная проверка "
            f"сняла ещё {RESEARCH.final}: {FINAL_REASONS} → по одному лиду на компанию: {s.leads_before} (ещё "
            f"{RESEARCH.spare} — вторые контакты тех же компаний, {RESEARCH.held} сняты вручную) → сборка базы "
            f"заново запросила каждую страницу: в базе все {s.total}.")


def stale_documents(s: SimpleNamespace, lines: dict[str, str]) -> tuple[list[str], list[str]]:
    """Numbers of README.md and of the hand-in notes that no longer match the data, and the documents checked.

    The hand-in notes are not part of the repository, so the file may be absent.
    """
    rows_word = plural(s.import_rows, "строка", "строки", "строк")
    expected = {
        README_MD: [
            *lines.values(),
            LEAD_RULE,
            funnel_line(s),
            f"Первая волна — приоритет A: {', '.join(s.first_wave)}.",
            LEADFINDER_LINE,
            *s.no_answer,  # the sites that did not answer on the last check are named in the limitations
        ],
        HANDIN_NOTES: [
            f"{s.total} лидов",
            f"именных адресов — {s.personal}, ящик должности ЛПР — {s.kinds[SERVICE_BOX]}",
            f"{s.import_rows} {rows_word} из {s.all_rows}",
            f"{s.their_swapped} строках из {s.their_total}",
            f"резерв — {s.reserve}",
        ],
    }
    problems, checked = [], []
    for name, phrases in expected.items():
        path = ROOT / name
        if not path.exists():
            continue
        checked.append(name)
        text = path.read_text(encoding="utf-8")
        problems += [f"{name}: нет строки «{phrase}»" for phrase in phrases if phrase not in text]
        if name == README_MD and not s.no_answer and README_NO_ANSWER in text:
            problems.append(f"{name}: абзац «…{README_NO_ANSWER}…» устарел — в базе нет строк, "
                            "чья страница не ответила")
    return problems, checked


# ---------- drawing ----------

def status_fill(value: str) -> PatternFill | None:
    if value.startswith(("OK",)):
        return FILL_OK
    if value.startswith(("РАСХОЖДЕНИЕ", "ПЕРЕПУТАНО")):
        return FILL_BAD
    if value.startswith("ПРОВЕРИТЬ"):
        return FILL_REVIEW
    return None


def set_widths(ws, widths: list[int]) -> None:
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width


def write_table(ws, headers: list[str], rows: list[list], top: int = 1, status_cols: tuple[int, ...] = (),
                autofilter: bool = True) -> int:
    """Header row at `top` + data rows, wrapped text; returns the number of the last row."""
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=top, column=col, value=header)
        cell.font, cell.fill, cell.alignment, cell.border = HEADER_FONT, HEADER_FILL, WRAP_TOP, BORDER
    for offset, row in enumerate(rows, 1):
        for col, value in enumerate(row, 1):
            cell = ws.cell(row=top + offset, column=col, value=value)
            cell.font, cell.alignment, cell.border = BODY_FONT, WRAP_TOP, BORDER
            if col - 1 in status_cols:
                fill = status_fill(str(value or ""))
                if fill:
                    cell.fill = fill
    last = top + len(rows)
    if autofilter:
        ws.auto_filter.ref = f"A{top}:{get_column_letter(len(headers))}{last}"
    return last


def text_height(text: str, width: float) -> float:
    """Row height in points for wrapped text in a column `width` wide (one character per width unit, with a margin)."""
    per_line = max(1, int(width))
    lines = sum(max(1, math.ceil(len(part) / per_line)) for part in str(text).split("\n"))
    return max(15.0, 13.5 * lines + 3)


def wide_row(ws, row: int, text: str, font: Font, first_col: int = 1, last_col: int = 1) -> None:
    """One line of text across merged columns. Merged cells do not auto-fit, so the height is set here."""
    cell = ws.cell(row=row, column=first_col, value=text)
    cell.font, cell.alignment = font, WRAP_TOP
    if last_col > first_col:
        ws.merge_cells(start_row=row, start_column=first_col, end_row=row, end_column=last_col)
    width = sum(ws.column_dimensions[get_column_letter(col)].width for col in range(first_col, last_col + 1))
    # a row may hold two merged blocks side by side: keep the taller one
    ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 0, text_height(text, width))


# ---------- sheets ----------

def base_table(ws, rows: list[dict], lead_columns: list[tuple[str, str, int]]) -> None:
    """Rows of task1_2_enriched.csv (or of the reserve, the same columns) as a sheet.

    A column that is empty in every row is not shown: a lead has no department mailbox behind it.
    """
    columns = [*lead_columns, *BASE_COLUMNS]
    used = {column for column, _, _ in columns + BASE_TAIL}
    if used != set(rows[0]):
        raise SystemExit(f"лист «{ws.title}» и его CSV расходятся в колонках: {sorted(used ^ set(rows[0]))}")
    columns = [(column, header, width) for column, header, width in columns if any(row[column] for row in rows)]
    headers = [header for _, header, _ in columns] + ["Вычитка"] + [header for _, header, _ in BASE_TAIL]
    data = [[row[column] for column, _, _ in columns] + [review_status(row["Комментарий"])]
            + [row[column] for column, _, _ in BASE_TAIL] for row in rows]
    set_widths(ws, [width for _, _, width in columns] + [18] + [width for _, _, width in BASE_TAIL])
    # The coloured status column is looked up by header, so adding columns cannot shift it.
    write_table(ws, headers, data, status_cols=(headers.index("Проверка_соответствия"),))
    ws.freeze_panes = "A2"  # the leading columns are too wide to freeze as well


def sheet_base(wb: Workbook, rows: list[dict]) -> None:
    ws = wb.create_sheet(SHEET_BASE)
    base_table(ws, rows, [])
    if [ws.cell(row=1, column=col).value for col in range(1, 7)] != COURSE_HEADERS:
        raise SystemExit(f"первые шесть колонок листа «{SHEET_BASE}» должны быть {COURSE_HEADERS}")


def sheet_reserve(wb: Workbook, rows: list[dict]) -> None:
    """Companies with a department mailbox only: they are not leads and are not launched."""
    ws = wb.create_sheet(SHEET_RESERVE)
    base_table(ws, rows, [RESERVE_REASON])


def chain_notes(s: SimpleNamespace) -> list[tuple[str, Font]]:
    def title(text: str) -> tuple[str, Font]:
        return text, SECTION

    def note(text: str) -> tuple[str, Font]:
        return text, BODY_FONT

    return [
        title("Кому какой текст"),
        note(f"База состоит из лидов, поэтому основной текст письма 1 — вариант «{NAMED_VARIANT}»: человеку, по имени. "
             "Письма 2 и 3 тоже начинаются с имени. Варианты «б» (ЛПР назван, адрес — ящик отдела) и «а» (ЛПР не "
             f"назван) — запасные: для листа «{SHEET_RESERVE}» и для вашей базы на листе «{SHEET_THEIR}». "
             f"В базе: {listing(s.variants, VARIANTS, quote=True)}."),
        title("Переменные и колонки базы"),
        note("{{companyName}} — колонка «{{companyName}} в письме»: короткое название бренда без «ООО» и кавычек, "
             "до 3 слов. Стоит только в начале темы, перед двоеточием, поэтому всегда в именительном падеже."),
        note("{{персонализация}} — колонка «Персонализация» листов 1-2 и 4: одно предложение до 30 слов от лица "
             "агентства, «Увидели, что вы…» или «Увидели на сайте, что вы…»."),
        note("{{гипотеза}} — колонка «Гипотеза_боли»: одно предложение до 16 слов сразу после персонализации. "
             "Всё, чего нет на сайте компании, написано как предположение."),
        note("{{firstName}} — колонка «Имя», а если сайт даёт отчество — «Имя» и «Отчество». Приветствие писем 1–3."),
        note("{{lprName}} — имя и фамилия, как они напечатаны на сайте компании; {{jobTitle}} — колонка "
             "«{{jobTitle}} в письме». Только запасной вариант «б»."),
        note("{{RANDOM | Здравствуйте | Добрый день}} — приветствие запасных вариантов «а» и «б»: сервис рассылок "
             "подставляет каждому получателю один из вариантов. В письмах 2 и 3 эта строка заменяет приветствие "
             f"по имени: она стоит в колонке «{GREETING_HEADER}»."),
        note("[Имя], [Имя Фамилия], [должность] — отправитель; заполняет команда агентства."),
        note("Если в «Персонализации» стоит «нет данных», письмо 1 этой компании не отправляем, пока не найден факт."),
        note("Слова считает tools/check_chain.py строгим счётом wc -w, тема входит в лимит. Худший случай "
             f"(персонализация 30 слов, гипотеза 16, название из 3 слов, имя с отчеством): тема + тело — не больше "
             f"{s.max_words} слов."),
        title("Логика прогрева"),
        note("Письмо 1 (день 1): факт о компании и гипотеза о её задаче. Один вопрос: привлечение новых клиентов — "
             "это вопрос к вам или к кому-то из коллег? В запасных вариантах: «а» — кто отвечает за новых клиентов "
             "и как с ним связаться; «б» — как связаться с названным ЛПР напрямую. Отказ — одной строкой: "
             "«Неактуально — ответьте «нет».»"),
        note("Письмо 2 (день 4, ответом в ту же ветку): тестовая неделя по шагам и соцдоказательство. "
             "CTA — ответить «да» и получить время для короткого созвона. Последняя строка просит переслать письмо "
             "коллеге, который отвечает за новых клиентов: должность не названа, потому что адресат может сам "
             "её занимать."),
        note("Письмо 3 (день 9): вежливый выход и польза — три вопроса для самопроверки перед запуском рассылок. "
             "CTA — «да» или «позже»."),
        note("Каждое письмо добавляет новое (факт и гипотеза → процесс → три вопроса), давление от письма к письму "
             "снижается. Цифры в письмах — только из курса: тестовая неделя, 7 дней, 300+ кейсов."),
        note(f"Если ответил человек — стоп по домену; автоответ не считается; см. лист «{SHEET_REPLIES}»."),
        note("Таблица «курс → цепочка», варианты для A/B и версия на 4 письма — в task3_chain.md."),
    ]


def sheet_chain(wb: Workbook, chain_rows: list[dict], enriched: list[dict], reserve: list[dict],
                contacts: list[dict], s: SimpleNamespace) -> None:
    ws = wb.create_sheet(SHEET_CHAIN)
    headers = ["Шаг", "Вариант письма 1", "Когда", "Тема", "Тема для A/B", "Текст письма", GREETING_HEADER,
               "Слов в шаблоне (тело с подписью)", "Макс. слов: тема + тело с подстановкой"]
    labels = {NAMED_VARIANT: f"«{NAMED_VARIANT}» — основной: лиду, по имени"}
    data = [[int(r["шаг"]), labels.get(r["вариант"], f"«{r['вариант']}» — запасной") if r["вариант"] else "все",
             r["задержка"], r["тема"], r["тема_AB"], r["текст"],
             r[GREETING_COLUMN], int(r["слов"]), int(r["слов_тема_плюс_тело_макс"])] for r in chain_rows]
    widths = [6, 16, 22, 32, 28, 90, 24, 14, 16]
    set_widths(ws, widths)
    last = write_table(ws, headers, data)
    ws.freeze_panes = "A2"
    for number, row in enumerate(data, 2):
        ws.row_dimensions[number].height = max(text_height(row[5], widths[5]), text_height(row[3], widths[3]))

    first, wide = 4, 6  # notes and examples span «Тема» … «Текст письма»
    line = last + 2
    for text, font in chain_notes(s):
        wide_row(ws, line, text, font, first, wide)
        line += 1

    # Rendered examples of email 1, through the renderer of the export: no raw variable can be left.
    chain = load_chain()
    sources = {OWN: enriched, RESERVE: reserve}
    line += 1
    wide_row(ws, line, "Примеры письма 1 с подстановкой: три лида, строка резерва и строка вашей базы", SECTION,
             first, wide)
    for contact, _ in preview_examples(contacts):
        letter = render_chain(contact, chain)[0]
        if "{{" in letter.subject + letter.body:
            raise SystemExit(f"{contact['company']}: в примере осталась переменная")
        proof = f"Факт: {contact['source']}."
        if contact["lprName"]:
            proof += f" Имя и должность: {sources[contact['base']][contact['row'] - 1]['источник_имени']}."
        variant = contact["letter1_variant"]
        why = THEIR_NOTE if contact["base"] == THEIR else VARIANT_NOTES[variant]
        if contact["address_type"] == SERVICE_BOX:
            why += f"; {ROLE_BOX_NOTE}"
        kind = {OWN: "Лид", RESERVE: f"Резерв, запасной вариант «{variant}»",
                THEIR: f"Ваша база, запасной вариант «{variant}»"}[contact["base"]]
        words = f"Слов: {letter.words} (тема {letter.subject_words} + тело {letter.body_words}), строгий счёт."
        block = [
            (f"{kind} — {contact['companyName']}: {why}.", BOLD),
            (f"Кому: {contact['email']}", BODY_FONT),
            (f"Тема: {letter.subject}", BOLD),
            (letter.body, BODY_FONT),
            (f"{words} {proof}", SMALL),
        ]
        line += 1
        for text, font in block:
            line += 1
            wide_row(ws, line, text, font, first, wide)


def sheet_their_base(wb: Workbook, their: list[dict], s: SimpleNamespace, lines: dict[str, str]) -> None:
    ws = wb.create_sheet(SHEET_THEIR)
    script = {r["company"]: r for r in read_csv("task4_script_output.csv")}
    corrected = {r["company"]: r for r in read_csv("task4_corrected_output.csv")}
    headers = [header for _, header, _ in THEIR_COLUMNS] + ["вердикт скрипта по исходной строке",
                                                           "вердикт скрипта после исправления"]
    data = [[r[column] for column, _, _ in THEIR_COLUMNS]
            + [script[r["company"]]["Проверка_соответствия"], corrected[r["company"]]["Проверка_соответствия"]]
            for r in their]
    set_widths(ws, [width for _, _, width in THEIR_COLUMNS] + [44, 36])
    last = write_table(ws, headers, data, status_cols=(headers.index("проверка"), len(headers) - 2, len(headers) - 1))
    ws.freeze_panes = "B2"
    notes = [
        "Номера строк в тексте — как на этом листе: строка 1 — заголовок, первая компания — строка 2.",
        ("«проверка» и исправления — ручной ресёрч: сайт и контакты каждой компании, MX домена, карточки в каталоге "
         "выставки «Металлообработка-2024» (Экспоцентр). Письма не отправлялись, SMTP-проверок не было."),
        ("«вердикт скрипта» — что сам personalize.py сказал по исходной строке и по исправленной. Все 6 подмен скрипт "
         "пометил РАСХОЖДЕНИЕ; ПРОВЕРИТЬ — где проверить автоматически нельзя (название иероглифами, сайт отдаёт "
         "скрипту 403)."),
        ("Персонализация — под слот {{персонализация}} письма 1, до 30 слов, «Увидели…»; каждый факт перепроверен "
         "03.10.2026 на странице из колонки «источник». Подробно: task4_traps.md."),
        (f"Сегмент, язык письма, триггер и гипотеза боли размечены сверх ТЗ. {lines['their']} Строки с языком EN "
         "в запуск не идут: цепочка написана на русском. Людей в этой базе я не искал: письмо идёт на адрес "
         "компании запасным вариантом «а», без обращения по имени."),
        ("База, судя по всему, собрана из каталога выставки 2024 года; 4 адреса из 15 текущие сайты не подтверждают. "
         "Курс называет порог Bounce Rate 5% — такой список нельзя запускать без валидатора. Это оценка риска, "
         "а не измеренный bounce."),
    ]
    for offset, text in enumerate(notes):
        cell = ws.cell(row=last + 2 + offset, column=1, value=text)
        cell.font, cell.alignment = SMALL, Alignment(wrap_text=False)


def sheet_replies(wb: Workbook, replies: list[dict]) -> None:
    ws = wb.create_sheet(SHEET_REPLIES)
    fields = [table for table in md_tables(REPLY_MD) if table[0].startswith(REPLY_FIELDS_SECTION)]
    if len(fields) != 1:
        raise SystemExit(f"{REPLY_MD.name}: таблица раздела «{REPLY_FIELDS_SECTION}» не найдена")
    heading, field_headers, field_rows = fields[0]
    set_widths(ws, [width for _, _, width in REPLY_COLUMNS])
    span = len(REPLY_COLUMNS) - 3  # intro lines and the second table end at «Текст шаблона»
    wide_row(ws, 1, first_line(REPLY_MD), BOLD, 1, span)
    wide_row(ws, 2, "Урок 4 курса: типы ответов, сроки и готовые шаблоны. Всё в квадратных скобках […] заполняет "
                    "человек в момент ответа или команда Polza: цены, кейсы и слоты я не выдумываю. Те же строки — "
                    "в reply_playbook.csv, полный текст — в reply_playbook.md.", BODY_FONT, 1, span)
    top = 4
    last = write_table(ws, [header for _, header, _ in REPLY_COLUMNS],
                       [[row[column] for column, _, _ in REPLY_COLUMNS] for row in replies], top=top)
    ws.freeze_panes = f"A{top + 1}"

    line = last + 2
    wide_row(ws, line, f"{heading} (reply_playbook.md, раздел {REPLY_FIELDS_SECTION.rstrip('.')})", SECTION, 1, span)
    line += 1
    split = 3  # «Поле» takes the first three columns, «Значения» the rest of the span
    for row in [field_headers, *field_rows]:
        header = row is field_headers
        for text, first, last_col in ((row[0], 1, split), (" · ".join(row[1:]), split + 1, span)):
            wide_row(ws, line, text, HEADER_FONT if header else BODY_FONT, first, last_col)
            for col in range(first, last_col + 1):
                ws.cell(row=line, column=col).border = BORDER
                if header:
                    ws.cell(row=line, column=col).fill = HEADER_FILL
        line += 1
    wide_row(ws, line + 1, "В reply_playbook.md также: скорость ответа, квалификация и передача отделу продаж, "
                           "когда цепочка останавливается, план Б при молчании, nurture.", SMALL, 1, span)


def sheet_launch(wb: Workbook, launch: list, s: SimpleNamespace) -> None:
    ws = wb.create_sheet(SHEET_LAUNCH)
    set_widths(ws, LAUNCH_WIDTHS)
    span = len(LAUNCH_WIDTHS)
    wide_row(ws, 1, first_line(LAUNCH_MD), BOLD, 1, span)
    wide_row(ws, 2, "Уроки 1 и 5 курса: четыре таблицы из launch_plan.md. Нормы — из курса, с уроком и пунктом; "
                    "числа в штуках — мой пересчёт на эту базу, а не прогноз результата.", BODY_FONT, 1, span)
    scope = (f"Календарь и метрики посчитаны на верхнюю границу — {s.all_rows} контакт"
             f"{plural(s.all_rows, '', 'а', 'ов')}. В файле импорта сейчас {s.import_rows} "
             f"{plural(s.import_rows, 'строка', 'строки', 'строк')}, это {s.letters} "
             f"{plural(s.letters, 'письмо', 'письма', 'писем')}.")
    if s.excluded_language:
        scope += (f" Строки с языком письма EN (их {s.excluded_language}) в запуск не идут, пока цепочка есть "
                  "только на русском.")
    if s.excluded_other:
        scope += f" Строки без персонализации (их {s.excluded_other}) тоже не идут: письмо 1 без факта не уходит."
    wide_row(ws, 3, scope, BODY_FONT, 1, span)
    line = 4
    for heading, headers, rows in launch:
        line += 1
        wide_row(ws, line, heading, SECTION, 1, span)
        line = write_table(ws, headers, rows, top=line + 1, autofilter=False) + 1
    wide_row(ws, line + 1, "В launch_plan.md также: объём и инфраструктура, прогрев, стоп-правила, A/B темы письма 1, "
                           "отчётность и журнал отправки, инструменты.", SMALL, 1, span)


def readme_lines(s: SimpleNamespace, n: dict[str, str]) -> list[tuple[str, str]]:
    def h(text: str) -> tuple[str, str]:
        return "h", text

    def p(text: str) -> tuple[str, str]:
        return "p", text

    types = plural(s.reply_types, "тип", "типа", "типов")
    steps = plural(s.checklist_steps, "шага", "шагов", "шагов")
    role_boxes = s.kinds[SERVICE_BOX]
    test_batches, later_batches = ", ".join(TEST_BATCHES), " и ".join(LATER_BATCHES)
    phones = ("Личных телефонов и личных адресов нет: телефон в базе — номер отдела или общий номер компании, как он "
              "напечатан на сайте. У части номеров есть добавочный: он напечатан в карточке самого лида или в строке "
              "его отдела.")
    if s.mobile_code:
        phones += (f" У {len(s.mobile_code)} {plural(len(s.mobile_code), 'компании', 'компаний', 'компаний')} "
                   f"({', '.join(s.mobile_code)}) номер стоит на мобильном коде: это номер из шапки сайта, а не "
                   "из карточки человека.")
    phones += (" В резерве на мобильном коде стоят номера PALLETOPTOM и Radist.Online: они напечатаны в шапке "
               "или в подвале каждой страницы сайта.")
    lines = [
        ("title", "Polza Agency — тестовое «Вайбкодер-аутричер». Максим Лутан"),
        h("Мини-курс Polza → что в сдаче"),
        p("ICP (урок 2, п. 4.1) → блок «ICP: кого искал» ниже."),
        p(f"Прямые контакты ЛПР, а не info@ (урок 2, п. 2.1 и п. 3.2) → лист «{SHEET_BASE}»: каждая строка — лид "
          "с именем, должностью и рабочим адресом; «Источник имени» — страница сайта компании, где они напечатаны."),
        p("Валидация по уровням (введение, п. 4) → колонка «Валидация» листа 1-2 и первый шаг чек-листа на листе "
          f"«{SHEET_LAUNCH}»."),
        p("Сегменты и триггеры (урок 2, п. 5 и п. 7.4) → колонки «Сегмент ЛПР», «Вертикаль», «Регион», "
          "«Часовой пояс», «Тип триггера», «Приоритет» листа 1-2."),
        p("Гипотеза боли (урок 2, п. 5.3) → колонка «Гипотеза_боли» листов 1-2 и 4, переменная {{гипотеза}} письма 1."),
        p(f"Цепочка (урок 3) → лист «{SHEET_CHAIN}»."),
        p(f"Инфраструктура, прогрев, разгон (урок 1) → лист «{SHEET_LAUNCH}» и launch_plan.md."),
        p(f"Ответы (урок 4) → лист «{SHEET_REPLIES}» и reply_playbook.md."),
        p(f"Метрики, A/B, масштабирование (урок 5) → лист «{SHEET_LAUNCH}» и launch_plan.md."),
        p("Экспорт для сервиса рассылок (урок 2, п. 4.4) → папка export/ рядом со скриптом."),
        p("Письма не отправлялись."),
        h("Что где"),
        p(f"{SHEET_BASE} — задания 1 и 2: {s.total} лидов. Первые шесть колонок — как просит курс: «Имя», "
          "«Фамилия», «Должность», «Email», «Телефон», «Компания». Дальше — где напечатаны имя и адрес, тип адреса "
          "и его проверка, сегменты и триггеры, персонализация со ссылкой на источник, гипотеза боли, проверка "
          "строки и пометка вычитки."),
        p(f"{SHEET_CHAIN} — задание 3: три письма от Polza Agency, все три обращаются к человеку по имени; "
          "переменные, логика прогрева, примеры с подстановкой. Два запасных варианта письма 1 — для строк "
          "без прямого адреса ЛПР."),
        p(f"{SHEET_THEIR} — задание 4: ваши {s.their_total} компаний, найденные подмены, исправления, "
          "персонализация, сегмент, язык письма и гипотеза."),
        p(f"{SHEET_REPLIES} — урок 4 курса: {s.reply_types} {types} ответов с готовыми шаблонами и поля, которые "
          "фиксируем по каждому ответу."),
        p(f"{SHEET_LAUNCH} — уроки 1 и 5: чек-лист до старта, календарь отправки, метрики с порогами, шаги "
          "масштабирования."),
        p(f"{SHEET_RESERVE} — {s.reserve} компаний первой версии базы: у них на сайте есть только ящик отдела. "
          "Это не лиды, в запуск они не идут; колонка «Почему в резерве» называет причину."),
        p("Рядом со скриптом лежат: personalize.py и тесты, сборка и проверка базы, проверка цепочки, экспорт "
          "для сервиса рассылок (export/), launch_plan.md и reply_playbook.md целиком, leadfinder.py — поиск лидов "
          "на сайтах компаний. README.md — запуск, карта «курс → сдача», цифры базы и ограничения; DESIGN.md — как "
          "устроен скрипт персонализации; LEADFINDER.md — как устроен поиск лидов."),
        h("Подход"),
        p("База — это лиды, а не ящики отделов: курс называет поиск info@ на сайтах тратой времени с низкой "
          "конверсией (урок 2, п. 2.1). Первая версия базы состояла из ящиков sales@ и info@; я её отложил в резерв "
          "и собрал базу заново — из людей."),
        p("Каждую клетку можно проверить: у лида есть страница сайта его компании, где имя, должность и адрес "
          "напечатаны рядом; у сигнала продаж — цитата и URL; у персонализации — URL и дословная цитата-основание; "
          "у каждой нормы в плане запуска и плейбуке — урок и пункт курса."),
        p("Сначала скрипты и валидаторы, потом ручная вычитка их результата. Готовую работу я отдавал на независимую "
          "проверку (отдельные агенты искали ошибки) и исправлял найденное: гонку TLS в потоках, слабые и "
          "устаревшие факты, завышенные роли контактов, склонение названия компании в письмах, страницы, закрытые "
          "в robots.txt."),
        h("ICP: кого искал"),
        p(f"Отрасль — {len(s.verticals)} вертикалей в трёх группах: IT и интеграторы, услуги для бизнеса, "
          "промышленность. Список с числами — в блоке «Задание 1»."),
        p("Размер — малый и средний бизнес, который продаёт компаниям. Признак на сайте: отдел продаж или "
          "коммерческий директор в контактах, форма «Запросить КП», дилерская программа, вакансия в продажи. "
          "Число сотрудников сайты почти не публикуют, поэтому отдельной колонки нет; где сайт его называет "
          "и оно близко к границе, это записано в «Примечании»."),
        p("Гео — Россия."),
        p("Кто считается ЛПР: первое лицо (собственник, генеральный, исполнительный или управляющий директор, "
          "управляющий партнёр), коммерческий директор, директор или руководитель отдела продаж, руководитель "
          "маркетинга или развития. Менеджеры, секретари, бухгалтерия, HR и техподдержка — нет."),
        p("Исключения: сервисы рассылок, лидогенерации и email-агентства (конкуренты Polza), крупные компании, "
          "сайты без свежих признаков жизни. Последнее правило применялось так: кандидат снимался, если самая "
          "поздняя дата на всём сайте (новости, карта сайта, подвал) старше 24 месяцев. Сайт, где лента новостей "
          "остановилась раньше, но есть признак моложе 24 месяцев (страница с контактом менялась позже, свежая "
          "страница называет того же человека, сайт обслуживается и в подвале стоит текущий год), оставался "
          "в базе с пометкой в «Примечании»."),
        p(funnel_line(s)),
        h("Задание 1. База"),
        p(f"{n['leads']} Дублей по домену и адресу нет, пересечений с вашей базой нет."),
        p(LEAD_RULE),
        p(f"{n['roles']} Один человек на компанию; где сайт называет нескольких, взят ближайший к продажам, "
          "остальные записаны как запасные контакты и в базу не пошли."),
        p(n["verticals"]),
        p(n["geo"]),
        p("Почему эти компании: на сайте каждой есть признак, что она продаёт B2B. Цитата и URL стоят в колонке "
          "«Сигнал». Конкурентов Polza и гигантов в базе нет; пограничные случаи (смежный рынок, верхняя граница "
          "среднего бизнеса, дочерняя структура) названы в «Примечании»."),
        p(f"{n['triggers']} Каждое событие в день сборки перепроверено по странице: вакансия — по названию на "
          "странице вакансий, дилеры — по слову на странице, выставка и расширение — по названию на странице, "
          "с которой взят факт."),
        p(f"{n['priority']} A — есть событие и ЛПР отвечает за продажи; B — есть событие или адрес именной; "
          f"C — ящик должности без события. Первая волна — приоритет A, {len(s.first_wave)} "
          f"{plural(len(s.first_wave), 'строка', 'строки', 'строк')}: {', '.join(s.first_wave)}."),
        p(n["reserve"] + f" Они лежат на листе «{SHEET_RESERVE}» и в task1_reserve.csv."),
        h("Лиды: что беру и откуда"),
        p("Источник — только сайт самой компании: «Контакты», «Команда», «Руководство», страницы отделов, свои "
          "новости. Каталоги и поиск использовались, только чтобы найти компанию и страницу. Из LinkedIn, hh.ru, "
          "справочников юрлиц, баз-брокеров, соцсетей и СМИ не взят ни один контакт."),
        p("Инструменты соблюдают robots.txt. Перед проверкой адрес приводится к виду, в котором уйдёт на сервер "
          "(без «/./», «/../» и их записи через «%2E»); с правилами сверяется и запрашивается одна и та же запись. "
          "Файл хоста читается до его первой страницы. С ним сверяется запрашиваемый адрес и цель каждой "
          "переадресации (их не больше 5 подряд); закрытая страница не запрашивается. Ключа, который отключает "
          "проверку, в командной строке нет. Лид, чья страница с контактом закрыта, в базу не идёт."),
        p("Что значит ответ на /robots.txt, решает один код для всех инструментов. Ответ 2xx, где есть хотя бы одна "
          "директива, — это правила. 4xx, кроме 429, и ответ 2xx без директив (пустой, HTML-страница) — файла нет, "
          "ничего не закрыто. 429, 5xx, таймаут, нет соединения, переадресация самого файла, которая к файлу "
          "не привела, — правила неизвестны, и страницы этого хоста не запрашиваются, пока файл не прочитан."),
        p("Чьи правила действуют. personalize.py и leadfinder.py называют себя в User-Agent (OutreachResearchBot, "
          "LeadFinderBot) и подчиняются группе, которая называет ровно это имя, а если её нет — группе "
          "«User-agent: *». Инструменты сборки базы запрашивают страницы через curl с обычным браузерным User-Agent "
          "и подчиняются только «User-agent: *». Прочитанный robots.txt personalize.py и leadfinder.py хранят "
          "в кэше сутки, сборка базы спрашивает его один раз за запуск."),
        p("Границы. Это поведение закреплено тестами: 55 ответов на /robots.txt, на каждом три загрузчика обязаны "
          "принять одно решение, и случаи из DESIGN.md. Ответ сервера, которого среди них нет, может быть обработан "
          "иначе. Группу с именем одного робота соблюдает только он; Crawl-delay не учитывается; адрес сравнивается "
          "с правилом буквально («//» и «%2F» — обычные знаки пути)."),
        p("Записано только то, что нужно для письма: имя, должность, рабочий адрес, страница-источник и её дата, "
          f"если она указана. {phones}"),
        p("Адрес по шаблону не подбирался и не достраивался, даже когда шаблон очевиден. SMTP-запросов к чужим "
          "почтовым серверам не было."),
        p(f"Именной адрес и ящик должности — разные вещи. Ящик должности ({role_boxes} "
          f"{plural(role_boxes, 'строка', 'строки', 'строк')}: dir@, ceo@, kd@, director@, rop@ или номерной ящик "
          "в карточке этого же человека) может разбирать помощник; такие строки помечены в колонке «Тип адреса» "
          "и без события идут в последнем батче."),
        p("Вторая проверка каждого кандидата была независимой: страница запрашивалась заново, расстояние от фамилии "
          f"до адреса измерялось по тексту. Отклонено {RESEARCH.rejected} — то, что формально находилось, но лидом "
          "не является: адрес дальше 350 знаков от имени, ящик отдела в карточке руководителя (marketing@, fin@), "
          "автор статьи в блоге, слишком крупная компания, сайт без свежих признаков жизни, страница, закрытая "
          "в robots.txt, сайт с запретом на обработку опубликованных персональных данных, компания, которая "
          "просит не присылать предложения."),
        p(f"Финальная проверка перед сдачей сняла ещё {RESEARCH.final}: {FINAL_REASONS}. Вручную сняты "
          f"{RESEARCH.held}: дочерняя структура крупной группы и адрес, который состоит из номера мобильного "
          "телефона."),
        p("В бою источник лидов — DealRocket и валидатор агентства: так эту задачу решает курс (урок 2, п. 3). "
          "В тестовом база собрана с сайтов компаний, чтобы каждую строку можно было проверить по ссылке. Сайты "
          "отдают прямой адрес ЛПР редко; замеры — в LEADFINDER.md."),
        h("Как проверены email"),
        p("Курс называет три уровня проверки адреса: синтаксис, MX-запись, актуальность ящика (введение, п. 4)."),
        p(f"1) Синтаксис — сделано: {s.total} из {s.total}."),
        p("2) MX-запись — сделано: есть у каждого домена."),
        p("3) Живость ящика — не проверялась. SMTP-проверку (RCPT TO) сознательно не делал: это зондирование чужого "
          f"почтового сервера. Шаг стоит первым в чек-листе на листе «{SHEET_LAUNCH}», его делает сервис-валидатор."),
        p("Сверх этих уровней: адрес есть в HTML страницы из колонки «Источник имени» (сравнение идёт после "
          "раскодирования HTML-сущностей), напечатан видимым текстом и стоит не дальше 350 знаков от фамилии; "
          "домен адреса — домен сайта или тот, на котором сайт печатает и другие свои "
          "адреса; это не общий ящик; robots.txt сайта эту страницу не закрывает. Проверка повторяется при каждой "
          "сборке: tools/build_base_all.py и tools/enrich_base.py запрашивают страницы заново."),
        p(f"Итог по строке — в колонке «Валидация», дата — в «Дата проверки» ({', '.join(s.checked)})."),
        p("Порог курса: при Bounce Rate выше 5% базу нужно чистить (введение, п. 5), лучше держать его ниже 2% "
          "(урок 5, п. 2.3). Базу перепроверяю перед запуском и пересобираю через 3 месяца: курс называет "
          "обновление данных «каждые 3 месяца» (урок 2, п. 3.4)."),
    ]
    if s.no_answer:
        lines.append(p(
            f"При последней сборке ({', '.join(s.checked)}) не ответила страница с контактом у "
            f"{len(s.no_answer)} {plural(len(s.no_answer), 'строки', 'строк', 'строк')}: {', '.join(s.no_answer)}. "
            "Адрес оставлен по проверке при сборке базы, строка помечена в «Примечании» и в «Валидации»; перед "
            "запуском она проверяется заново и снимается, если страница не поднимется."))
    lines += [
        h("Задание 2. Скрипт персонализации"),
        p("personalize.py: на входе CSV (company, site, email), на выходе те же строки плюс колонки "
          "«Персонализация», «Источник», «Проверка_соответствия», «Комментарий». LLM-шаг — Claude Code в "
          "headless-режиме (claude -p, модель sonnet)."),
        p("Как не выдумать факт: модель возвращает URL и дословную цитату, а код проверяет:"),
        p("— страница с этим URL реально загружена, и цитата на ней есть;"),
        p("— все числа текста есть в цитате;"),
        p("— текст на русском, это одно предложение до 30 слов, начинается с «Увидели…»;"),
        p("— нет слов «недавно» и «в этом году» без свежей даты."),
        p("Если проверка не прошла, модель отвечает ещё раз, зная причину отказа. Если и тогда не прошла, в "
          "колонку пишется «нет данных», а кандидат уходит в комментарий."),
        p("Свежесть: скрипт читает страницы новостей, релизов и кейсов и поднимает самые новые датированные записи. "
          "В промпте указана сегодняшняя дата. Факт старше 12 месяцев помечается в комментарии."),
        p("Скрипт работает с данными уровня компании: людей он не ищет. Лиды собраны отдельно; адрес строки скрипт "
          "только сверяет со страницей из колонки email_source. Проверки кодом на имена людей в тексте нет: "
          "в трёх строках из 98 скрипт поставил в персонализацию имя человека, который не адресат письма. Вычитка "
          "это убрала. В task2_script_output.csv такие имена заменены на «[имя]» в четырёх строках: в тексте "
          "скрипта и в цитате — у этих трёх, только в цитате — у четвёртой. В итоговом файле «[имя]» стоит "
          "в комментариях 5 строк, в колонке «Персонализация» его нет."),
        p(f"Результат скрипта — task2_script_output.csv: {s.script_ok} OK и {s.script_other} ПРОВЕРИТЬ, "
          f"{s.script_done} персонализаций и {s.script_empty} «нет данных». Скрипт шёл двумя прогонами: 206 с "
          "на 69 компаний прежнего состава базы и 165 с на 35 строк (34 новые и один повтор). Затем я вычитал все "
          f"{s.total} строк по страницам-источникам: {s.review_kept} оставил дословно, в {s.review_reworded} оставил "
          f"факт скрипта и поправил формулировку, {s.review_replaced} заменил другим фактом. Чаще всего скрипт брал "
          "самую свежую справочную статью блога вместо факта о самой компании. У каждой правки в комментарии есть "
          "причина, цитата и вариант скрипта. На листе 1-2 итог вычитки, фильтр — колонка «Вычитка»."),
        p("Гипотеза боли (колонка «Гипотеза_боли») — не LLM, а словарь в tools/enrich_base.py: формулировка "
          "на вертикаль. Если у компании подтверждён триггер (вакансия в продажи, набор дилеров), формулировка "
          "берётся по нему; «между выставками» говорится только строкам с триггером «выставка». Всё, чего нет "
          "на сайте компании, написано как предположение."),
        h("Задание 3. Цепочка"),
        p("Три письма по ТЗ, собраны по уроку 3 курса. Представляется человек («Меня зовут [Имя], я из Polza "
          "Agency»), действия агентства описаны на «мы», в подписи — имя, фамилия и должность."),
        p("1) знакомство: обращение по имени, факт о компании, гипотеза боли и один вопрос — привлечение новых "
          "клиентов — это к вам или к кому-то из коллег?;"),
        p("2) через 3 дня, ответом в ту же ветку: тестовая неделя по шагам и почему рабочая почта клиента "
          "не участвует;"),
        p("3) через 5 дней: вежливый выход и три вопроса для самопроверки."),
        p("Все три письма начинаются с имени адресата. Для строк без прямого адреса ЛПР (резерв, ваша база) "
          "у письма 1 есть два запасных варианта: «а» — кто отвечает за новых клиентов; «б» — на сайте указан ЛПР, "
          "как связаться напрямую; письма 2 и 3 им идут с безличным приветствием."),
        p("В каждом письме один CTA. Цифры — только из курса: тестовая неделя, 7 дней, 300+ кейсов. "
          "{{companyName}} стоит только в начале темы, поэтому название не склоняется. Тема и тело вместе — "
          f"не больше {s.max_words} слов, если каждая переменная взята по своему максимуму из task3_chain.md; "
          "с самыми длинными значениями из базы каждое письмо тоже укладывается в лимит ТЗ, 120 слов. Считает "
          "tools/check_chain.py."),
        p("В task3_chain.md — таблица «курс → цепочка», варианты для A/B и приложение с версией на 4 письма, "
          "как рекомендует курс."),
        h("Задание 4. Ваша база: ловушки"),
        p(f"Подмены нашлись в {s.their_swapped} строках из {s.their_total} (строки 4, 5, 6, 7, 8 и 12 на листе "
          f"«{SHEET_THEIR}»):"),
        p("— чужие сайты: Ункомтех у Tengzhong, Юнимаш у Tesid;"),
        p("— у ТД Ункомтех email и сайт двух других китайских компаний;"),
        p("— названия JAT и Rogen переставлены местами;"),
        p("— один сайт saintymachine.com указан у двух компаний, а у Shixinghong ещё и бесплатный ящик 126.com."),
        p("Ещё 2 строки похожи на ошибку, но верны (Internor, HNC). У 4 строк адреса нет на текущем сайте: для трёх "
          "предложена замена, адрес Fengyi подтверждён каталогами 2024 и 2026 годов."),
        p("Как нашёл: открыл сайт и контакты каждой компании, проверил MX домена и сверил все 15 строк с карточками "
          "в каталоге выставки «Металлообработка-2024» на сайте Экспоцентра. База, судя по всему, собрана оттуда."),
        p("Скрипт на исходной базе помечает РАСХОЖДЕНИЕ ровно эти 6 строк и подсказывает правильный домен, если "
          "он есть в другой строке. После исправлений — 11 OK, 4 ПРОВЕРИТЬ, 0 РАСХОЖДЕНИЕ. Повторный прогон "
          "03.10.2026 дал те же статусы."),
        p(f"Сверх ТЗ размечены сегмент, язык письма, триггер и гипотеза боли. {n['their']} Строки с языком EN "
          "в запуск не идут: цепочка написана на русском."),
        p("Каталог 2024 года — старый источник: 4 адреса из 15 текущие сайты не подтверждают. Курс называет порог "
          "Bounce Rate 5% — такой список нельзя запускать без валидатора. Это оценка риска, а не измеренный bounce."),
        h("Сверх ТЗ: план запуска, ответы, экспорт, поиск лидов"),
        p("План запуска (уроки 1 и 5): отдельный домен и 2 ящика, прогрев не меньше 14 дней, чек-лист из "
          f"{s.checklist_steps} {steps}, календарь по дням, метрики с порогами тревоги, стоп-правила, A/B темы "
          "письма 1, шаги масштабирования. Нормы — из курса; числа в штуках — мой пересчёт на эту базу, "
          "а не прогноз."),
        p(f"Плейбук ответов (урок 4): {s.reply_types} {types} ответов, сроки реакции, шаблоны, статусы лида, "
          "квалификация, план Б при молчании. Цены, слоты и кейсы в шаблонах — заглушки […]: их даёт команда Polza."),
        p(f"{n['export']} В тесте — батчи {test_batches}, то есть письма 1 первой недели; батчи {later_batches} "
          f"({s.waiting} {plural(s.waiting, 'лид', 'лида', 'лидов')}) уходят после замера и получают "
          f"тему-победителя; батч {THEIR_BATCH} — ваша база, запасной вариант «а». Файл export/import_instantly.csv "
          "собирает tools/export_campaign.py; кода отправки в нём нет. Письмо 1 с подстановкой для каждой строки — "
          "в export/letters_preview.csv, примеры для чтения вслух — в export/preview.md."),
        p("leadfinder.py — инструмент, который делает тот же поиск автоматически: по ICP находит компании, обходит "
          "их сайты и выдаёт ЛПР с адресом и страницей-доказательством, а где лида нет — причину. "
          f"{LEADFINDER_LINE} Размеченный набор собран 03.10.2026; в нём остался один контакт, который 04.10.2026 "
          "снят из базы как устаревший, и цифры посчитаны вместе с ним. Как инструмент устроен, как мерился "
          "и чем ограничен — в LEADFINDER.md."),
        p("Ничего не отправлялось: домены не покупались, сервисы не подключались."),
        h("Ограничения и что сделал бы дальше"),
        p("Живость ящиков не проверялась: это делает валидатор перед запуском. Адрес на сайте мог устареть — "
          "человек сменил работу, а страницу не обновили. Где страница сама говорит о возрасте (старые новости, "
          "документ с датой), это записано в «Примечании»."),
        p(f"{s.total} лидов — это то, что сайты отдают честно; масштаб в бою дают DealRocket и валидатор агентства, "
          "как в курсе (урок 2, п. 3). "
          f"{role_boxes} {plural(role_boxes, 'адрес', 'адреса', 'адресов')} из {s.total} — ящики должности: "
          "письмо обращается по имени, но прочитать его может помощник."),
        p("Цена правила про robots.txt: два лида сняты, потому что страница с контактом закрыта; у одной компании "
          "факт взят не из закрытого раздела новостей, а со страницы «О компании». Сайт, чей robots.txt в день "
          "запуска не получен (5xx, 429, таймаут, нет соединения), в этом запуске не обходится: в персонализации "
          "строка получает «нет данных» и берётся заново при следующем запуске, в сборке базы такая страница "
          "считается не ответившей."),
        p("Что в работе с robots.txt осталось разным у инструментов: попыток получить файл — до двух у "
          "personalize.py, одна или две у leadfinder.py, три у curl; первые два читают не больше 3 МБ файла "
          "и держат его в кэше сутки, curl читает целиком и не хранит; сайт, закрывшийся от одного из роботов "
          "по имени, остальные инструменты читают. Директива с опечаткой пропускается. Подробно — в README.md, "
          "раздел «Ограничения»."),
        p(f"Цепочка есть только на русском: строки вашей базы с языком письма EN (их {s.their_languages['EN']}) "
          "ждут английской версии."),
        p("Гипотеза боли — формулировка на вертикаль или на триггер. Это предположение, а не знание о компании; "
          "проверяется ответами на тестовой неделе."),
        *([p(f"У {len(s.future_facts)} {plural(len(s.future_facts), 'лида', 'лидов', 'лидов')} "
             f"({', '.join(s.future_facts)}) факт — событие, которое в день сборки ещё впереди: выставка, форум или "
             "вебинар октября и ноября 2026 года. Они помечены в «Примечании»: после даты события формулировка "
             "переводится в прошедшее время.")] if s.future_facts else []),
        p("jimmytool.com и jat-carbide.com в прогоне 03.10.2026 ответили скрипту 403. Скрипт не маскируется под "
          "браузер, поэтому факты для листа 4 по ним сверены в тот же день отдельным запросом curl. 04.10.2026 оба "
          "сайта отвечают кодом 202 и страницей капчи — и скрипту, и браузерному User-Agent."),
        p("Сайты, которые рисуются только JavaScript, скрипт не читает: headless-браузера нет."),
        p("Китайские названия, записанные иероглифами, с латиницей не сопоставляются, такие строки получают "
          "ПРОВЕРИТЬ."),
        p("Скрипт предпочитает самую свежую запись и часто выбирает справочную статью. Следующий шаг: отличать "
          "материалы о самой компании (кейс, релиз, выставка) от справочных статей. Пока это ловит вычитка."),
        p(f"Дальше: см. листы «{SHEET_LAUNCH}» и «{SHEET_REPLIES}». Тему оцениваю по OR на тестовой неделе, "
          "подтверждаю по RR."),
        h("Найдено и исправлено по ходу"),
        p("— Стандартный парсер robots.txt в Python ложно запрещал 4 сайта («Disallow: /?» читается как "
          "«Disallow: /»)."),
        p("— Сборка базы не спрашивала robots.txt. Теперь читает файл до первой страницы хоста и сверяет с ним "
          "адрес перед запросом; две строки с закрытой страницей контактов сняты, один факт заменён."),
        p("— Скрипт персонализации шёл за переадресацией, не сверив адрес назначения с robots.txt, хранил "
          "robots.txt в кэше бессрочно и имел ключ, отключающий проверку. Теперь сверяется каждый шаг "
          "переадресации, файл живёт в кэше сутки, страница из кэша сверяется с действующими правилами, "
          "неполученный файл закрывает хост до следующего запуска, ключа нет."),
        p("— Метка порядка байтов в начале robots.txt прятала первую группу правил, а правило с кириллицей "
          "не совпадало с процентной записью того же пути."),
        p("— Адрес с «/../» сверялся с robots.txt в одной записи, а на сервер уходил в другой: «/a/../private/x» "
          "проходил правило «Disallow: /private/», а запрашивался «/private/x». Теперь адрес приводится к одной "
          "записи до проверки, и запрашивается она же."),
        p("— Разборщик правил был в двух копиях, а ответ на /robots.txt каждый загрузчик читал своим кодом. "
          "Теперь правила, разбор ответа и запись адреса — один код в personalize.py; leadfinder.py и сборка базы "
          "берут его оттуда. Таблица из 55 ответов требует от трёх загрузчиков одного решения."),
        p("— leadfinder.py не шёл за переадресацией самого robots.txt. Теперь идёт, в том числе на другой сайт, "
          "и применяет правила найденного файла к хосту, который спрашивал. Параметр respect_robots у его "
          "загрузчика убран."),
        p("— curl разворачивал «{a,b}» и «[1-3]» в адресе в список адресов и запрашивал каждый. Теперь он "
          "запускается с ключом -g, а с ключом -q не читает ~/.curlrc."),
        p("— Сайт с неполной цепочкой сертификатов не открывался."),
        p("— Общий SSL-контекст у потоков иногда пропускал чужой сертификат."),
        p("— Адреса в HTML-сущностях не находились."),
        p("— Баннер cookies в <article> прятал весь текст страницы."),
        p("— При продолжении прерванного прогона комментарии с кавычками внутри обрезались."),
        p("Для каждого случая есть тест."),
        h("Задание 5"),
        p("Текст про вайбкод-стек написал сам, без ИИ. Он в отдельном документе, ссылка — в сообщении."),
    ]
    return lines


def sheet_readme(wb: Workbook, lines: list[tuple[str, str]]) -> None:
    ws = wb.create_sheet(SHEET_README)
    ws.column_dimensions["A"].width = 130
    for kind, text in lines:
        if kind == "h":
            ws.append([])
        ws.append([text])
        cell = ws.cell(row=ws.max_row, column=1)
        cell.font = {"title": TITLE, "h": SECTION}.get(kind, BODY_FONT)
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def main(argv: list[str] | None = None) -> int:
    argv = sys.argv[1:] if argv is None else argv
    enriched, their = read_csv("task1_2_enriched.csv"), read_csv("task4_final.csv")
    reserve, leads_before = read_csv("task1_reserve.csv"), len(read_csv("task1_leads.csv"))
    chain_rows, replies = read_csv("task3_chain.csv"), read_csv("reply_playbook.csv")
    launch = md_tables(LAUNCH_MD)
    if len(launch) != LAUNCH_TABLES:
        raise SystemExit(f"{LAUNCH_MD.name}: таблиц {len(launch)}, а лист «{SHEET_LAUNCH}» ждёт {LAUNCH_TABLES}")
    contacts, report = build_contacts()
    stats = collect_stats(enriched, their, contacts, report, chain_rows, launch, replies, reserve, leads_before,
                          read_csv("task2_script_output.csv"))
    numbers = number_lines(stats)
    if "--numbers" in argv:
        print("\n".join([*numbers.values(), funnel_line(stats)]))
        return 0

    wb = Workbook()
    wb.remove(wb.active)
    sheet_base(wb, enriched)
    sheet_chain(wb, chain_rows, enriched, reserve, contacts, stats)
    sheet_their_base(wb, their, stats, numbers)
    sheet_replies(wb, replies)
    sheet_launch(wb, launch, stats)
    sheet_reserve(wb, reserve)
    sheet_readme(wb, readme_lines(stats, numbers))
    wb.save(OUT)
    print(f"saved {OUT}")
    print(f"листов {len(wb.sheetnames)}: " + "; ".join(wb.sheetnames))
    print(f"строк: база {len(enriched)}, резерв {len(reserve)}, цепочка {len(chain_rows)}, ваша база {len(their)}, "
          f"ответы {len(replies)}, таблицы плана запуска {stats.launch_rows}")
    print("числа листа README:")
    print("\n".join("  " + line for line in [*numbers.values(), funnel_line(stats)]))

    stale, checked = stale_documents(stats, numbers)
    if stale:
        print("\nДОКУМЕНТЫ УСТАРЕЛИ (таблица собрана; исправьте текст и запустите ещё раз):")
        print("\n".join("  - " + problem for problem in stale))
        return 1
    print(f"OK: числа в {' и '.join(checked)} совпадают с базой")
    return 0


if __name__ == "__main__":
    sys.exit(main())
